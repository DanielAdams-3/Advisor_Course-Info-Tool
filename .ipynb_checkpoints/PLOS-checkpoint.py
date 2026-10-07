#import pandas as pd

#import "C:\Users\daad2295\AppData\Local\Python\pythoncore-3.14-64\Lib\site-packages\openpyxl"
import openpyxl
#how to check PYTHONPATH? this program runs fine on my desktop but I can't run it in VS Code on my desktop?
#-------------------------------------------------------------------------------------------------------------------------------------------------------
#Define Student class
#-------------------------------------------------------------------------------------------------------------------------------------------------------
class Student:
    stu_info_keys = []
    stu_info_values = []
    stu_info_dict = {}
    plos_file_name = ""
    Slate_file = ""

    def __init__(self, plos_file_name="template",Slate_file="Plan of Study Templater.xlsx",stu_info_keys=["CU-SID","Person Name","Pronouns","Person Visa Type","Study Plan Codes","Subplan Code 1","Subplan Code 2","Admit Term", "Person Visa Type", "Transfer Courses","BAM Supplement","Student Classes", "Waivers", "Overall GPA", "Other Study Plans"]):
        self.Slate_file = Slate_file
        self.stu_info_keys = stu_info_keys
        self.plos_file_name=plos_file_name
#-------------------------------------------------------------------------------------------------------------------------------------------------------
#Function - get # of plans of study
#-------------------------------------------------------------------------------------------------------------------------------------------------------
def get_num_plos(num_plos):
    #access and open Slate file
    #path="C:\\Users\\daad2295\\Desktop\\PLOS_Maker\\Plan of Study Templater.xlsx"
    path="R:\\gradadmin\\Grad Program\\3. Daniel Adams_GPS\\Student Forms\\Plan of Study\\PLOS_Maker\\Plan of Study Templater.xlsx"
    #path='Plan of Study Templater.xlsx'
    slate_wb=openpyxl.load_workbook(path)
    slate_sheet=slate_wb.active

    #identify number of plos to create, minus the header row
    num_plos = slate_sheet.max_row - 1
    return num_plos
#-------------------------------------------------------------------------------------------------------------------------------------------------------
#Function - import data from Slate to PLOS.py
#-------------------------------------------------------------------------------------------------------------------------------------------------------
def import_from_slate(plos_list, row_num):
    #access and open Slate file
    path="R:\\gradadmin\\Grad Program\\3. Daniel Adams_GPS\\Student Forms\\Plan of Study\\PLOS_Maker\\Plan of Study Templater.xlsx"
    #path='Plan of Study Templater.xlsx'
    slate_wb=openpyxl.load_workbook(path)
    slate_sheet=slate_wb.active
        
    current = Student() #variables should be template at this point
    c=1
    r = row_num #user_provided
    current.stu_info_values.clear()
    for key in current.stu_info_keys:
        cell_obj=slate_sheet.cell(r,c)
        current.stu_info_values.append(cell_obj.value)
        c +=1
    current.stu_info_dict = dict(zip(current.stu_info_keys,current.stu_info_values))
    #print(current.stu_info_dict)
    #update the path variable here, so we can differentiate between plans of study by program
    if ("CSEN-MSCPS" in current.stu_info_dict["Study Plan Codes"]):
        #current.plos_file_name="C:\\Users\\daad2295\\Desktop\\PLOS_Maker\\Plan of Study_MSCPS.xlsx"
        current.plos_file_name="R:\\gradadmin\\Grad Program\\3. Daniel Adams_GPS\\Student Forms\\Plan of Study\\PLOS_Maker\\Plan of Study_MSCPS.xlsx"
        #current.plos_file_name='Plan of Study_MSCPS.xlsx'
    elif("NTEN-MSNE" in current.stu_info_dict["Study Plan Codes"]):
        #current.plos_file_name="C:\\Users\\daad2295\\Desktop\\PLOS_Maker\\Plan of Study_MSNE.xlsx"
        current.plos_file_name="R:\\gradadmin\\Grad Program\\3. Daniel Adams_GPS\\Student Forms\\Plan of Study\\PLOS_Maker\\Plan of Study_MSNE.xlsx"
        #current.plos_file_name='Plan of Study_MSNE.xlsx'
    elif("AINT-MSAI" in current.stu_info_dict["Study Plan Codes"]):
        current.plos_file_name="R:\\gradadmin\\Grad Program\\3. Daniel Adams_GPS\\Student Forms\\Plan of Study\\PLOS_Maker\\Plan of Study_MSAIP.xlsx"
        #current.plos_file_name='Plan of Study_MSAIP.xlsx'
    plos_list.append(current)
    #print(plos_list)

#-------------------------------------------------------------------------------------------------------------------------------------------------------
#actual program start 
#-------------------------------------------------------------------------------------------------------------------------------------------------------

plos_list = []
num_plos = 0
num_plos = get_num_plos(num_plos)
row_num = 1 

for i in range (num_plos):
    row_num += 1
    import_from_slate(plos_list, row_num)

for stu in plos_list:
    path = stu.plos_file_name
    #print(path)
    plos_wb=openpyxl.load_workbook(path)
    plos_sheet=plos_wb['Data']

    #save as new file with new name
    #new_file_path = "C:\\Users\\daad2295\\Desktop\\PLOS_Maker\\"
    new_file_path="R:\\gradadmin\\Grad Program\\3. Daniel Adams_GPS\\Student Forms\\Plan of Study\\"

    if ("CSEN-MSCPS" in stu.stu_info_dict["Study Plan Codes"]):
        if (stu.stu_info_dict["Subplan Code 1"] != None):
            if(stu.stu_info_dict["Subplan Code 2"] != None):
                new_file_name = "Plos_" + stu.stu_info_dict["Person Name"] + "_MSCPS, " + stu.stu_info_dict["Subplan Code 1"] + ", " + stu.stu_info_dict["Subplan Code 2"] + ".xlsx"
            else:
                new_file_name = "Plos_" + stu.stu_info_dict["Person Name"] + "_MSCPS, " + stu.stu_info_dict["Subplan Code 1"] + ".xlsx"
        elif (stu.stu_info_dict["Subplan Code 2"] != None):
            new_file_name = "Plos_" + stu.stu_info_dict["Person Name"] + "_MSCPS-" + stu.stu_info_dict["Subplan Code 2"] + ".xlsx"
        else:
            new_file_name = "Plos_" + stu.stu_info_dict["Person Name"] + "_MSCPS.xlsx"
            
    elif ("NTEN-MSNE" in stu.stu_info_dict["Study Plan Codes"]):
        if (stu.stu_info_dict["Subplan Code 1"] != None):
            new_file_name = "Plos_" + stu.stu_info_dict["Person Name"] + "_MSNE, " + stu.stu_info_dict["Subplan Code 1"] + ".xlsx"
        elif (stu.stu_info_dict["Subplan Code 2"] != None):
            new_file_name = "Plos_" + stu.stu_info_dict["Person Name"] + "_MSNE, " + stu.stu_info_dict["Subplan Code 2"] + ".xlsx"
        else:
            new_file_name = "Plos_" + stu.stu_info_dict["Person Name"] + "_MSNE.xlsx"
    '''
    elif ("AINT-MSAIP" in stu.stu_info_dict["Study Plan Codes"]):
        if (stu.stu_info_dict["Subplan Code 1"] != None):
            new_file_name = "Plos_" + stu.stu_info_dict["Person Name"] + "_MSAIP, " + stu.stu_info_dict["Subplan Code 1"] + ".xlsx"
        elif (stu.stu_info_dict["Subplan Code 2"] != None):
            new_file_name = "Plos_" + stu.stu_info_dict["Person Name"] + "_MSAIP, " + stu.stu_info_dict["Subplan Code 2"] + ".xlsx"
        else:
            new_file_name = "Plos_" + stu.stu_info_dict["Person Name"] + "_MSAIP.xlsx"
    '''
    #print(new_file_name)
    new_file_path += new_file_name
    plos_wb.save(new_file_path)

   #fill it with Student Information
    e = plos_sheet['D2']
    e.value = stu.stu_info_dict["CU-SID"]
    e = plos_sheet['D3']
    e.value = stu.stu_info_dict["Person Name"]
    e = plos_sheet['D4']
    e.value = stu.stu_info_dict["Pronouns"]
    e = plos_sheet['D5']
    e.value = stu.stu_info_dict["Study Plan Codes"]
    e = plos_sheet['D6']
    e.value = stu.stu_info_dict["Subplan Code 1"]
    e = plos_sheet['D7']
    e.value = stu.stu_info_dict["Subplan Code 2"]
    e = plos_sheet['D8']
    e.value = stu.stu_info_dict["Admit Term"]
    e = plos_sheet['D9']
    e.value = stu.stu_info_dict["Person Visa Type"]
    e = plos_sheet['D10']
    e.value = stu.stu_info_dict["Waivers"]
    e = plos_sheet['D11']
    e.value = stu.stu_info_dict["Overall GPA"]
    e = plos_sheet['D13']
    e.value = stu.stu_info_dict["Other Study Plans"]
        
    #process Student Classes column
    student_classes=[]
    temp = ""
    temp = stu.stu_info_dict["Student Classes"]
    if (temp != None):
        student_classes=temp.split('*')

    #process Transfer Classes column
    transfer_classes=[]
    temp = ""
    temp = stu.stu_info_dict["Transfer Courses"]
    if (temp != None):
        transfer_classes = temp.split('*')
        index=0
        for transfer_class in range(len(transfer_classes)):
            student_classes.append(transfer_classes[index])
            index+=1
        
    #process BAM Supplement column ONLY IF TRANSFER CLASSES WAS EMPTY AND DIDN'T HAVE 'BAM' IN IT.
    bam_classes=[]
    temp = ""
    temp = stu.stu_info_dict["BAM Supplement"]
    if (stu.stu_info_dict["Transfer Courses"] == None):
        if (temp != None):
            bam_index=0
            bam_classes = temp.split('*')
            for bam_class in range(len(bam_classes)):
                student_classes.append(bam_classes[bam_index])
                bam_index+=1

    #add classes to plos
    course_pieces={}
    r = 1
    c = 5
    index = 0
    course_piece = ""
    for course in range(len(student_classes)):
        course_pieces= student_classes[index].split(',')
        r +=1
        for p in range(len(course_pieces)):
            course_piece = course_pieces[p].strip()
            if (c == 5):
                if(len(course_piece)>8):
                    new=course_piece[0:4]
                    new = new + ' '
                    new = new + course_piece[5:]
                    course_piece = new
            cell_obj=plos_sheet.cell(r,c)
            cell_obj.value = course_piece
            c +=1
        c = 5
        index +=1

    #close out
    plos_wb.save(new_file_path)
    student_classes.clear()
    bam_classes.clear()
    transfer_classes.clear()
    course_pieces.clear()
